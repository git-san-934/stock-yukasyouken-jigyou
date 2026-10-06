import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

metrics, jpx_path, out, fetched = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]
m = pd.read_csv(metrics, dtype={"code": str})
markets = {}
try:
    j = pd.read_csv(jpx_path, dtype=str)
    markets = dict(zip(j["コード"].str.strip(), j["市場・商品区分"]))
except Exception as e:
    print("jpx list unavailable:", e)
m["market"] = m["code"].map(markets)
if markets:
    m = m[m["market"].fillna("").str.contains("プライム|スタンダード|グロース")]

COLS = [
    ("code", "銘柄コード", None, 10),
    ("name", "銘柄名", None, 34),
    ("market", "市場", None, 18),
    ("float_ratio", "浮動株割合", "0.0%", 12),
    ("bd_big", "大株主／全株式", "0.0%", 11),
    ("bd_fund", "投資信託等／全株式", "0.0%", 11),
    ("bd_foreign", "外国人／全株式", "0.0%", 11),
    ("bd_other", "その他／全株式", "0.0%", 11),
    ("fixed_B", "不動株割合（東証式近似）", "0.0%", 14),
    ("fixed_A", "特定株比率（四季報式）", "0.0%", 14),
    ("gross_margin", "粗利率", "0.0%", 10),
    ("cash_m", "保有現金（百万円）", "#,##0", 14),
    ("dps", "1株配当（円）", "#,##0.00", 10),
    ("yield_note", "利回りの注意", None, 22),
    ("div_yield", "配当利回り", "0.00%", 10),
    ("buyback_status", "自社株買の状況（直近期）", None, 46),
    ("buyback_m", "当期自己株式取得額（百万円）", "#,##0", 14),
    ("treasury_ratio", "自己株式比率", "0.0%", 10),
    ("officer_ratio", "役員持株比率", "0.0%", 10),
    ("top10_ratio", "上位10大株主比率", "0.0%", 10),
    ("top1", "筆頭株主", None, 40),
    ("period_end", "決算期末", None, 11),
    ("gm_basis", "粗利率の基準", None, 8),
    ("cash_basis", "現金の内訳", None, 18),
    ("payout", "配当性向", "0.0%", 9),
    ("price", "株価（円）", "#,##0.0", 10),
    ("price_date", "株価日付", None, 11),
    ("doc_id", "EDINET書類ID", None, 12),
]
m["cash_m"] = m["cash"] / 1e6
m["yield_note"] = None
if "split_adj" in m:
    m.loc[m["split_adj"].notna(), "yield_note"] = "決算期後の分割" + m["split_adj"].fillna(0).astype(int).astype(str) + "倍で調整"
if "yield_check" in m:
    m.loc[(m["yield_check"] == 1) & m["yield_note"].isna(), "yield_note"] = "8%超・分割/特別配当を要確認"
m["buyback_m"] = m["buyback_amount"] / 1e6


def sheet(wb, title, df, sort="fixed_B"):
    ws = wb.create_sheet(title)
    df = df.sort_values(sort)
    hdr = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1F4E78")
    for c, (_, label, _, w) in enumerate(COLS, 1):
        cell = ws.cell(1, c, label)
        cell.font, cell.fill = hdr, fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[1].height = 42
    for r, (_, row) in enumerate(df.iterrows(), 2):
        for c, (key, _, fmt, _) in enumerate(COLS, 1):
            v = row.get(key)
            if pd.isna(v):
                v = None
            cell = ws.cell(r, c, v)
            if fmt:
                cell.number_format = fmt
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLS))}{max(len(df) + 1, 1)}"
    return len(df)


wb = Workbook()
info = wb.active
info.title = "説明"
nF = sheet(wb, "抽出_浮動株5%未満", m[m["float_ratio"] < 0.05], "float_ratio")
nB = sheet(wb, "参考_不動株5%未満", m[m["fixed_B"] < 0.05])
nA = sheet(wb, "参考_特定株5%未満", m[m["fixed_A"] < 0.05], "fixed_A")
sheet(wb, "全銘柄", m, "code")
pmin = m["price_date"].dropna().min() if "price_date" in m else ""
pmax = m["price_date"].dropna().max() if "price_date" in m else ""
lines = [
    ("東証銘柄抽出：浮動株割合5%未満", True),
    (f"データ取得日：{fetched}", False),
    (f"対象：{len(m):,}社（東証プライム・スタンダード・グロース、直近の有価証券報告書がある会社）", False),
    (f"抽出結果：浮動株割合5%未満 {nF}社（参考：不動株割合（東証式近似）5%未満 {nB}社 ／ 特定株比率（四季報式）5%未満 {nA}社）", False),
    ("", False),
    ("■ 浮動株割合の定義", True),
    ("・浮動株割合 ＝（全株式 − 大株主 − 投資信託など金融機関 − 外国人 − その他の法人など）÷ 全株式", False),
    ("　有報「所有者別状況」の個人その他の割合（単元株ベース）から、自己株式、大株主上位10名のうちの個人、役員持株（上位10名に入らない分）を差し引いて計算。", False),
    ("　大株主の個人は名前（漢字・ひらがなの氏名で、法人を示す語がないもの）で判定。資産管理会社経由の保有は「その他の法人」側に入ります。", False),
    ("・内訳（全株式に対する割合、浮動株＋4項目で100%）", False),
    ("　大株主＝上位10大株主の合計。投資信託等＝所有者別状況の金融機関＋金融商品取引業者から、大株主に入っている金融機関の分を除いたもの。", False),
    ("　外国人＝外国法人等から大株主に入っている外国人の分を除いたもの。その他＝残り（その他の法人、自己株式、役員持株、政府など）。", False),
    ("　有報には投資信託だけの数字がないため、信託銀行・生保・銀行などの金融機関をまとめて「投資信託等」としています。", False),
    ("", False),
    ("■ 不動株割合の定義（参考、発行済株式総数に対する割合）", True),
    ("・不動株割合（東証式近似）＝ 10%以上を持つ大株主 ＋ 役員持株 ＋ 自己株式等", False),
    ("　信託銀行の信託口・海外カストディアン名義（マスタートラスト、日本カストディ銀行、STATE STREET等）は投資家の預かり株なので固定株に数えていません。", False),
    ("　東証の流通株式の定義では銀行・保険・事業法人の政策保有株も非流通ですが、有報からは判別できないため含めていません（実際の不動株はこれより高めになります）。", False),
    ("・特定株比率（四季報式）＝ 上位10大株主 ＋ 役員持株 ＋ 自己株式等（参考）", False),
    ("　上位10名には信託口も含まれるため、5%未満になる会社は通常ありません。", False),
    ("・役員と大株主の重複は、10位の株数以上を持つ役員（Bは10%以上の役員）を大株主側に含まれるとみなして除いています。", False),
    ("", False),
    ("■ 各項目", True),
    ("・粗利率＝売上総利益÷売上高（連結優先、連結がなければ単体）。銀行・保険など売上総利益がない業種は空欄。", False),
    ("・保有現金＝連結の現金及び預金（IFRS企業は現金及び現金同等物）。", False),
    (f"・配当利回り＝有報の1株当たり配当（直近期実績）÷株価（終値 {pmin}〜{pmax}）。決算期後の株式分割は有報のPER×EPSと現在株価の比から推定して調整。期中の分割・特別配当で高く出ることがあり、8%超は「利回りの注意」欄に印を付けています。", False),
    ("・自社株買の状況＝有報「自己株式の取得等の状況」で取締役会/株主総会決議による取得の記載有無と、キャッシュ・フロー計算書の自己株式の取得額（直近期）。", False),
    ("　有報提出後に発表された自社株買い（適時開示）は反映していません。", False),
    ("", False),
    ("■ 出典", True),
    ("・各社の有価証券報告書（EDINET API v2、XBRL CSV）：大株主の状況、役員の状況、自己株式等、経営指標等、財務諸表", False),
    ("・株価：Yahoo Finance（yfinance 経由）", False),
    ("・市場区分：日本取引所グループ「東証上場銘柄一覧」", False),
    ("・抽出スクリプト：github.com/git-san-934/stock-yukasyouken-jigyou（ブランチ claude/stock-screen）", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = info.cell(i, 1, t)
    if b:
        c.font = Font(bold=True, size=12 if i > 1 else 14)
info.column_dimensions["A"].width = 130
wb.save(out)
print("saved", out, "F:", nF, "B:", nB, "A:", nA, "total:", len(m))
